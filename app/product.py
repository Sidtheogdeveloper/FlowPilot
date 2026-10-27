"""Owner-scoped product tools. Existing authentication and tenancy remain authoritative."""
import csv,io,json,math,re,uuid
from datetime import timedelta
from typing import Literal
from fastapi import APIRouter,HTTPException,UploadFile,File,Form,Query
from fastapi.responses import Response
from pydantic import BaseModel,Field
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from . import models as db,artifacts
from .tenancy import owner

router=APIRouter(prefix='/product',tags=['Product tools'])
def identity():
    value=owner.get()
    if not value:raise HTTPException(401,'Sign in to continue')
    return value
def dataset(s,key):
    row=s.get(db.Dataset,key)
    if not row:raise HTTPException(404,'Dataset not found')
    return row
def workflow(s,key):
    row=s.get(db.Workflow,key)
    if not row:raise HTTPException(404,'Workflow not found')
    return row
def root_id(flow):return (flow.planner or {}).get('chat_id') or flow.id

def preferences(items,view='active'):
    user=owner.get()
    if not user:return items
    with db.Session() as s:settings={p.chat_id:p for p in s.scalars(select(db.ChatPreference).where(db.ChatPreference.user_id==user))}
    result=[]
    for item in items:
        p=settings.get(item['chat_id'])
        flags={'pinned':bool(p and p.pinned),'archived':bool(p and p.archived),'trashed':bool(p and p.trashed)}
        if (view=='trash' and not flags['trashed']) or (view!='trash' and flags['trashed']):continue
        if view=='archived' and not flags['archived'] or view=='active' and flags['archived']:continue
        result.append(dict(item,**flags,title=p.title if p and p.title else item['title']))
    return sorted(result,key=lambda x:not x['pinned'])

class ChatChange(BaseModel):
    title:str|None=Field(default=None,min_length=1,max_length=160)
    pinned:bool|None=None
    archived:bool|None=None

@router.patch('/chats/{key}')
def change_chat(key:str,payload:ChatChange):
    user=identity()
    with db.Session.begin() as s:
        root=root_id(workflow(s,key));s.get(db.Workflow,root,with_for_update=True)
        p=s.get(db.ChatPreference,(user,root))
        if not p:p=db.ChatPreference(user_id=user,chat_id=root);s.add(p)
        for field in payload.model_fields_set:
            value=getattr(payload,field)
            if value is not None:
                if field=='title':
                    value=value.strip()
                    if not value:raise HTTPException(422,'Enter a nonempty title')
                setattr(p,field,value)
    return {'ok':True}

@router.delete('/chats/{key}')
def trash_chat(key:str):
    user=identity()
    with db.Session.begin() as s:
        root=root_id(workflow(s,key));s.get(db.Workflow,root,with_for_update=True)
        s.execute(insert(db.ChatPreference).values(user_id=user,chat_id=root,trashed=True).on_conflict_do_update(index_elements=['user_id','chat_id'],set_={'trashed':True}))
    return {'ok':True,'message':'Moved to your trash. Restore it at any time.'}

@router.post('/chats/{key}/restore')
def restore_chat(key:str):
    user=identity()
    with db.Session.begin() as s:
        root=root_id(workflow(s,key));p=s.get(db.ChatPreference,(user,root),with_for_update=True)
        if p:p.trashed=False;p.archived=False
    return {'ok':True}

@router.get('/datasets/{key}/preview')
def preview(key:str,offset:int=Query(0,ge=0),limit:int=Query(25,ge=1,le=100)):
    from .tables import ensure_metadata
    with db.Session.begin() as s:
        d=dataset(s,key)
        if d.kind!='tabular':raise HTTPException(422,'Use the document source viewer for this file')
        profile=ensure_metadata(d);content=artifacts.read_json(d.table_ref)
        return {'id':d.id,'filename':d.filename,'total_rows':len(content['rows']),'offset':offset,'limit':limit,'columns':profile['columns'],'rows':content['rows'][offset:offset+limit],'profile':profile}

@router.get('/datasets/{key}/distribution')
def distribution(key:str,column:str,bins:int=Query(10,ge=2,le=30)):
    import pandas as pd
    from .tables import ensure_metadata
    with db.Session.begin() as s:
        d=dataset(s,key)
        if d.kind!='tabular':raise HTTPException(422,'A numeric table column is required')
        profile=ensure_metadata(d);meta=next((c for c in profile['columns'] if c['id']==column),None)
        if not meta or meta['type']!='number':raise HTTPException(422,'Select a resolved numeric column')
        values=pd.to_numeric(pd.Series([r[column] for r in artifacts.read_json(d.table_ref)['rows']]),errors='coerce').dropna()
        if not len(values):return {'points':[],'column':meta['name'],'total':0}
        counts=pd.cut(values,bins=bins,duplicates='drop',precision=4).value_counts(sort=False)
        return {'column':meta['name'],'total':len(values),'points':[{'label':str(k),'value':int(v)} for k,v in counts.items()],'scope':'Source dataset before analytical filters'}

def revision(s,d):
    record=s.get(db.DatasetRevision,d.id)
    if not record:
        record=db.DatasetRevision(dataset_id=d.id,family_id=d.id,revision=1);s.add(record);s.flush()
    return record

@router.get('/datasets/{key}/versions')
def versions(key:str):
    with db.Session.begin() as s:
        d=dataset(s,key);base=revision(s,d)
        return [{'id':item.id,'filename':item.filename,'revision':r.revision,'rows':item.rows,'created':r.created.isoformat(),'current':item.id==key} for r,item in s.execute(select(db.DatasetRevision,db.Dataset).join(db.Dataset,db.Dataset.id==db.DatasetRevision.dataset_id).where(db.DatasetRevision.family_id==base.family_id).order_by(db.DatasetRevision.revision.desc()))]

@router.post('/datasets/{key}/versions')
async def upload_version(key:str,file:UploadFile=File(...)):
    data=await file.read(25*1024*1024+1)
    if len(data)>25*1024*1024:raise HTTPException(413,'Maximum upload is 25 MiB')
    from .tables import materialize,ensure_metadata
    from .api import dataset_info
    with db.Session.begin() as s:
        d=dataset(s,key);base=revision(s,d)
        # Lock the family's original dataset to serialize revision numbering.
        s.get(db.Dataset,base.family_id,with_for_update=True)
        if d.kind!='tabular':raise HTTPException(422,'Document updates should be uploaded as a new document')
        try:frame,profile,table_ref,profile_ref=materialize(data,file.filename or 'update.csv',d.sheet)
        except (ValueError,TypeError) as exc:raise HTTPException(422,str(exc))
        old=ensure_metadata(d);old_ids={c['name']:c['id'] for c in old['columns']};used={int(c['id'][1:]) for c in old['columns']};next_id=max(used,default=-1)+1
        mapping={}
        for c in profile['columns']:
            new=old_ids.get(c['name'])
            if new is None:new=f'c{next_id}';next_id+=1
            mapping[c['id']]=new;c['id']=new
        content=artifacts.read_json(table_ref)
        content['schema']={mapping[k]:v for k,v in content['schema'].items()};content['rows']=[{mapping[k]:v for k,v in r.items()} for r in content['rows']]
        table_ref=artifacts.put_json(content)['id'];profile_ref=artifacts.put_json(profile)['id']
        artifact=artifacts.put(data,'.xlsx' if (file.filename or '').lower().endswith('.xlsx') else '.csv')['id']
        number=max(s.scalars(select(db.DatasetRevision.revision).where(db.DatasetRevision.family_id==base.family_id)),default=1)+1
        new=db.Dataset(id=str(uuid.uuid4()),filename=file.filename or d.filename,rows=len(frame),artifact=artifact,table_ref=table_ref,profile_ref=profile_ref,sheet=d.sheet)
        s.add(new);s.flush();s.add(db.DatasetRevision(dataset_id=new.id,family_id=base.family_id,revision=number))
        from .planning_tools import drift_between
        return dict(dataset_info(new),revision=number,drift=drift_between(d,new))

@router.get('/workflows/{key}/compare')
def compare(key:str,other:str|None=None):
    with db.Session() as s:
        current=workflow(s,key);prior_id=other or (current.planner or {}).get('parent_workflow_id')
        if not prior_id:return {'previous_id':None,'current_id':key,'added':current.plan['tasks'],'removed':[],'changed':[],'unchanged':[],'inputs_changed':False}
        prior=workflow(s,prior_id);old={t['id']:t for t in prior.plan['tasks']};new={t['id']:t for t in current.plan['tasks']}
        changed=[];same=[]
        for tid in new.keys()&old.keys():
            fields=[name for name in ('handler','parameters','inputs','depends_on','resources') if new[tid].get(name)!=old[tid].get(name)]
            (changed if fields else same).append({'id':tid,'fields':fields,'before':old[tid],'after':new[tid]} if fields else tid)
        return {'previous_id':prior.id,'current_id':key,'added':[new[t] for t in new.keys()-old.keys()],'removed':[old[t] for t in old.keys()-new.keys()],'changed':changed,'unchanged':sorted(same),'inputs_changed':prior.dataset_id!=current.dataset_id or (prior.planner or {}).get('dataset_ids')!=(current.planner or {}).get('dataset_ids')}

def accepted_tables(s,flow):
    tasks=list(s.scalars(select(db.TaskRow).where(db.TaskRow.workflow_id==flow.id,db.TaskRow.status=='completed')))
    # Prefer analytical outputs rather than loading the entire source table for chart menus.
    for t in tasks:
        if not t.artifact or t.artifact.endswith('.html') or t.spec['handler'] in ('table_load','table_load_chunk','table_partial','validate_dataset','clean_transactions','document_load','document_search','document_summarize'):continue
        content=artifacts.read_json(t.artifact)
        if isinstance(content,dict) and isinstance(content.get('rows'),list):yield t,content

@router.get('/workflows/{key}/charts')
def chart_data(key:str):
    with db.Session() as s:
        flow=workflow(s,key);tables=[]
        for task,data in accepted_tables(s,flow):
            rows=data['rows'];schema=data.get('schema') or {k:{'name':k,'type':'number' if isinstance(v,(int,float)) and not isinstance(v,bool) else 'text'} for k,v in (rows[0] if rows else {}).items()}
            tables.append({'id':task.task_id,'title':task.task_id.replace('_',' '),'schema':schema,'rows':rows[:250],'total_rows':len(rows),'audit':data.get('audit',[])})
        return {'tables':tables,'limit':250}

def spreadsheet_cell(value):
    # Text beginning with spreadsheet formula prefixes must remain text in exports.
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')):return "'"+value
    return value

@router.get('/workflows/{key}/exports/{format}')
def export_result(key:str,format:Literal['csv','xlsx','pdf'],task:str|None=None):
    from .api import detail
    item=detail(key)
    if not item.get('insights'):raise HTTPException(409,'Complete the analysis before exporting')
    if format=='pdf':
        from .report_export import make_pdf
        return Response(make_pdf(item),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="flowpilot-v{item["revision"]}.pdf"'})
    with db.Session() as s:
        tables=list(accepted_tables(s,workflow(s,key)))
    if task:tables=[x for x in tables if x[0].task_id==task]
    if not tables:raise HTTPException(422,'No accepted tabular result is available to export')
    if format=='csv':
        row,data=tables[0];columns=list(data.get('schema') or (data['rows'][0] if data['rows'] else {}));buffer=io.StringIO();writer=csv.writer(buffer)
        writer.writerow([spreadsheet_cell(data.get('schema',{}).get(k,{}).get('name',k)) for k in columns])
        writer.writerows([[spreadsheet_cell(r.get(k)) for k in columns] for r in data['rows']])
        return Response(('\ufeff'+buffer.getvalue()).encode('utf-8'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="flowpilot-result.csv"'})
    from openpyxl import Workbook
    book=Workbook(write_only=True)
    for i,(row,data) in enumerate(tables):
        sheet=book.create_sheet((re.sub(r'[\\/*?:\[\]]','_',row.task_id)[:25] or 'Result')+'_'+str(i+1))
        columns=list(data.get('schema') or (data['rows'][0] if data['rows'] else {}))
        sheet.append([spreadsheet_cell(data.get('schema',{}).get(k,{}).get('name',k)) for k in columns])
        for r in data['rows']:sheet.append([spreadsheet_cell(r.get(k)) for k in columns])
    buffer=io.BytesIO();book.save(buffer)
    return Response(buffer.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="flowpilot-results.xlsx"'})

class FeedbackRequest(BaseModel):
    category:Literal['correct','incorrect','inefficient','suggestion']
    text:str=Field(min_length=1,max_length=4000)
@router.post('/workflows/{key}/feedback',status_code=201)
def add_feedback(key:str,payload:FeedbackRequest):
    if not payload.text.strip():raise HTTPException(422,'Enter feedback text')
    with db.Session.begin() as s:
        workflow(s,key);value=db.Feedback(id=str(uuid.uuid4()),user_id=identity(),workflow_id=key,category=payload.category,text=payload.text.strip());s.add(value)
    return {'id':value.id,'message':'Feedback saved. It does not automatically retrain a model.'}
@router.get('/workflows/{key}/feedback')
def feedback(key:str):
    with db.Session() as s:
        workflow(s,key)
        return [{'id':f.id,'category':f.category,'text':f.text,'created':f.created.isoformat()} for f in s.scalars(select(db.Feedback).where(db.Feedback.workflow_id==key,db.Feedback.user_id==identity()).order_by(db.Feedback.created))]

class ScheduleRequest(BaseModel):
    workflow_id:str
    interval_minutes:int=Field(ge=15,le=525600)
    use_latest:bool=True
class RerunRequest(BaseModel):
    use_latest:bool=True
    reuse_outputs:bool=True
@router.post('/workflows/{key}/rerun')
def rerun(key:str,payload:RerunRequest):
    from .product_jobs import replay
    from .api import detail
    with db.Session.begin() as s:
        f=workflow(s,key)
        if f.status not in ('completed','failed','cancelled'):raise HTTPException(409,'Finish the current analysis first')
        try:result=replay(s,f,payload.use_latest,reuse=payload.reuse_outputs)
        except ValueError as exc:raise HTTPException(422,str(exc))
        new_id=result.id
    return detail(new_id)
@router.post('/schedules',status_code=201)
def schedule(payload:ScheduleRequest):
    with db.Session.begin() as s:
        f=workflow(s,payload.workflow_id)
        if f.status!='completed' or (f.planner or {}).get('kind') not in ('tabular','document'):raise HTTPException(422,'Schedule a completed table or document analysis')
        if len(list(s.scalars(select(db.Schedule).where(db.Schedule.owner_id==identity(),db.Schedule.active==True))))>=20:raise HTTPException(422,'Maximum 20 active schedules')
        item=db.Schedule(id=str(uuid.uuid4()),owner_id=identity(),workflow_id=f.id,interval_minutes=payload.interval_minutes,use_latest=payload.use_latest,next_run=db.now()+timedelta(minutes=payload.interval_minutes));s.add(item);s.flush()
        return {'id':item.id,'next_run':item.next_run.isoformat(),'notification_channel':'in-app'}
@router.get('/schedules')
def schedules():
    with db.Session() as s:
        return [{'id':x.id,'workflow_id':x.workflow_id,'interval_minutes':x.interval_minutes,'active':x.active,'use_latest':x.use_latest,'next_run':x.next_run.isoformat(),'error':x.error} for x in s.scalars(select(db.Schedule).where(db.Schedule.owner_id==identity()).order_by(db.Schedule.created.desc()))]
class ScheduleChange(BaseModel):active:bool
@router.patch('/schedules/{key}')
def change_schedule(key:str,payload:ScheduleChange):
    with db.Session.begin() as s:
        row=s.get(db.Schedule,key,with_for_update=True)
        if not row or row.owner_id!=identity():raise HTTPException(404,'Schedule not found')
        row.active=payload.active
        if payload.active:row.error=None;row.next_run=db.now()+timedelta(minutes=row.interval_minutes)
    return {'ok':True}
@router.get('/notifications')
def notifications():
    with db.Session() as s:
        return [{'id':n.id,'title':n.title,'body':n.body,'workflow_id':n.workflow_id,'read':n.read,'created':n.created.isoformat()} for n in s.scalars(select(db.Notification).where(db.Notification.owner_id==identity()).order_by(db.Notification.created.desc()).limit(100))]
@router.post('/notifications/{key}/read')
def read_notification(key:str):
    with db.Session.begin() as s:
        row=s.get(db.Notification,key)
        if not row or row.owner_id!=identity():raise HTTPException(404,'Notification not found')
        row.read=True
    return {'ok':True}
