"""Bounded ingestion and immutable, domain-independent table profiles."""
import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime
import pandas as pd
from . import artifacts

MAX_ROWS = 1_000_000
MAX_COLUMNS = 200
MAX_BYTES = 25 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024

def worksheets(data):
    if len(data) > MAX_BYTES:
        raise ValueError('Maximum upload is 25 MiB')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED:
            raise ValueError('Workbook exceeds 100 MiB decompressed limit')
        if len(archive.infolist()) > 10000:
            raise ValueError('Workbook contains too many entries')
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        return book.sheetnames
    finally:
        book.close()

def read_table(data, filename, sheet=None):
    if len(data) > MAX_BYTES:
        raise ValueError('Maximum upload is 25 MiB')
    if filename.lower().endswith('.xlsx'):
        names = worksheets(data)
        if sheet not in names:
            raise ValueError('Select an available worksheet')
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            iterator = book[sheet].iter_rows()
            headers = tuple(cell.value for cell in next(iterator, ()))
            if not 1 <= len(headers) <= MAX_COLUMNS:
                raise ValueError('Tables must contain 1–200 columns')
            rows = []
            for row in iterator:
                if len(row) > MAX_COLUMNS:
                    raise ValueError('Maximum 200 columns')
                if len(rows) >= MAX_ROWS:
                    raise ValueError('Dataset exceeds 1 million rows')
                values=[]
                for cell in row:
                    v=cell.value
                    # Excel can store identifiers numerically while displaying a
                    # fixed zero-padded format. Preserve those visible zeros too.
                    fmt=getattr(cell,'number_format','')
                    if isinstance(v,(int,float)) and not isinstance(v,bool) and float(v).is_integer() and re.fullmatch(r'0{2,}',fmt):
                        v=('-' if v<0 else '')+str(abs(int(v))).zfill(len(fmt))
                    values.append(v.isoformat() if isinstance(v,(date,datetime)) else v)
                rows.append(values)
        finally:
            book.close()
        frame = pd.DataFrame(rows, columns=headers)
        # Numeric Excel cells may be numbers; text identifiers remain strings.
        frame = frame.map(lambda v: '' if pd.isna(v) else str(v))
    elif filename.lower().endswith('.csv'):
        text = data.decode('utf-8-sig')
        reader = csv.reader(io.StringIO(text), strict=True)
        header = next(reader, [])
        if len(header) > MAX_COLUMNS:
            raise ValueError('Maximum 200 columns')
        # pandas otherwise silently renames duplicate headers.
        if not header or any(not h.strip() for h in header) or len(set(header)) != len(header):
            raise ValueError('Duplicate column names; rename them before upload')
        count = 0
        for row in reader:
            if not row:continue  # Same blank-line behavior as pandas.
            count += 1
            if count > MAX_ROWS:raise ValueError('Dataset exceeds 1 million rows')
            if len(row) != len(header):
                raise ValueError('CSV row width does not match its header at record '+str(count+1))
        frame = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, nrows=MAX_ROWS + 1)
    else:
        raise ValueError('Supported formats are CSV and .xlsx')
    if not 1 <= len(frame.columns) <= MAX_COLUMNS:
        raise ValueError('Tables must contain 1–200 columns')
    if any(not isinstance(h, str) or not h.strip() for h in frame.columns) or len(set(frame.columns)) != len(frame.columns):
        raise ValueError('Column names must be nonempty and unique')
    if len(frame) > MAX_ROWS:
        raise ValueError('Dataset exceeds 1 million rows')
    return frame.fillna('')

def profile(frame):
    columns = []
    for index, name in enumerate(frame.columns):
        values = frame[name].astype(str)
        present = values[values.str.strip() != '']
        numeric = pd.to_numeric(present, errors='coerce')
        finite = numeric.map(lambda v: pd.notna(v) and __import__('math').isfinite(float(v)))
        leading_zero = bool(present.str.match(r'^[+-]?0\d+$').any())
        iso = present.str.match(r'^\d{4}-\d{2}-\d{2}(?:[T ].*)?$')
        dates = pd.to_datetime(present.where(iso), errors='coerce', utc=True, format='mixed')
        ambiguous = int(present.str.match(r'^\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}$').sum())
        kind = 'text'
        if len(present) and finite.all() and not leading_zero:
            kind = 'number'
        elif len(present) and dates.notna().all():
            kind = 'datetime'
        elif len(present) and not leading_zero and (finite.any() or ambiguous):
            kind = 'unresolved'
        summary = None
        if kind == 'number':
            summary = {'min': float(numeric.min()), 'max': float(numeric.max()), 'mean': float(numeric.mean())}
        columns.append({'id': f'c{index}', 'name': name, 'type': kind,
                        'null_count': int((values.str.strip() == '').sum()),
                        'distinct_count': int(present.nunique()), 'numeric': summary,
                        'leading_zero_identifier': leading_zero,
                        'date_parse': {'iso_valid': int(dates.notna().sum()), 'ambiguous': ambiguous, 'nonempty': len(present)}})
    return {'version': 1, 'rows': len(frame), 'columns': columns}

def materialize(data, filename, sheet=None):
    frame = read_table(data, filename, sheet)
    description = profile(frame)
    content = {'rows': [{f'c{i}': str(value) for i, value in enumerate(row)} for row in frame.itertuples(index=False, name=None)],
               'schema': {c['id']: {'name': c['name'], 'type': c['type']} for c in description['columns']},
               'audit': []}
    table_ref = artifacts.put_json(content)['id']
    profile_ref = artifacts.put_json(description)['id']
    return frame, description, table_ref, profile_ref

def ensure_metadata(dataset):
    if dataset.profile_ref:
        return artifacts.read_json(dataset.profile_ref)
    _, description, table_ref, profile_ref = materialize(artifacts.read(dataset.artifact), dataset.filename)
    dataset.table_ref = table_ref
    dataset.profile_ref = profile_ref
    return description
